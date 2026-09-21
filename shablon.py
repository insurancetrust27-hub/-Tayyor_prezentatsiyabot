"""Tayyor Canva shablonlarini to'ldirish moduli.

Har bir shablon uchun bir marta "xarita" (SLOTLAR) yoziladi: qaysi slaydda,
qaysi matn qutisiga qanday matn tushishi va u necha belgidan oshmasligi kerak.
AI shu xaritaga qarab matn yozadi, kod esa shablonning dizaynini buzmasdan
faqat matnlarni almashtiradi.
"""
import io
import json
import os

from pptx import Presentation
from pptx.util import Pt

PAPKA = os.path.dirname(os.path.abspath(__file__))

# rejimlar:
#   "sarlavha" - katta bir qatorli sarlavha (KATTA HARFDA, sig'masa shrift kichrayadi)
#   "matn"     - oddiy matn (belgi limiti bor)
#   "katta"    - matn, lekin KATTA HARFDA yoziladi
#   "sobit"    - har doim bir xil matn (kalit kerak emas)
#   "tozala"   - bezak matnini bo'sh qoldiradi (Home / About / Contact)

def _s(slayd, shape_id, kalit, rejim, limit, tavsif):
    return {"slayd": slayd, "id": shape_id, "kalit": kalit,
            "rejim": rejim, "limit": limit, "tavsif": tavsif}


def _sobit(slayd, shape_id, matn):
    return {"slayd": slayd, "id": shape_id, "kalit": None,
            "rejim": "sobit", "limit": 0, "tavsif": "", "matn": matn}


def _tozala(slayd, shape_id):
    return {"slayd": slayd, "id": shape_id, "kalit": None,
            "rejim": "tozala", "limit": 0, "tavsif": ""}


MODA_SLOTLAR = [
    # 1-slayd: muqova
    _s(1, 23, "muqova_1", "sarlavha", 10, "Muqova sarlavhasining 1-qatori: mavzuning 1-so'zi yoki qisqa qismi"),
    _s(1, 22, "muqova_2", "sarlavha", 12, "Muqova sarlavhasining 2-qatori: mavzuning davomi"),
    _s(1, 24, "muqova_izoh", "matn", 42, "Muqova ostidagi bir qatorli qisqa shior"),
    _tozala(1, 25), _tozala(1, 26), _tozala(1, 27),
    _sobit(1, 28, "TAQDIMOT"),

    # 2-9: matnli slaydlar
    _s(2, 17, "s2_a", "matn", 420, "Kirish: mavzu nima ekanligi, asosiy g'oya (qalin ajratilgan matn)"),
    _s(2, 16, "s2_b", "matn", 460, "Kirishning davomi: mavzuning muhim jihatlari"),
    _s(3, 22, "s3_a", "matn", 335, "Mavzuning ahamiyati (qalin ajratilgan matn)"),
    _s(3, 18, "s3_b", "matn", 490, "Ahamiyatning davomi: misollar va tafsilotlar"),
    _s(4, 18, "s4_a", "matn", 310, "Mavzuning rivojlanishi yoki tarixi"),
    _s(4, 19, "s4_b", "matn", 520, "Bugungi holati va asosiy yo'nalishlar (qalin ajratilgan matn)"),
    _s(5, 14, "s5_a", "matn", 285, "Amaliy qo'llanilishi (qalin ajratilgan matn)"),
    _s(5, 13, "s5_b", "matn", 585, "Amaliy misollar va foydalari"),
    _s(6, 17, "s6", "katta", 295, "Asosiy xususiyatlar yoki ko'nikmalar haqida bitta abzats"),
    _s(7, 17, "s7_a", "matn", 285, "Muammolar va qiyinchiliklar"),
    _s(7, 18, "s7_b", "matn", 550, "Qiyinchiliklarni yengish yo'llari (qalin ajratilgan matn)"),
    _s(8, 15, "s8", "katta", 305, "Afzalliklar va ijobiy tomonlari haqida bitta abzats"),
    _s(9, 15, "s9_a", "matn", 470, "Asosiy atamalar va ularning izohi (2-3 atama)"),
    _s(9, 16, "s9_b", "matn", 370, "Yana 2-3 ta atama va izohi"),

    # 10: 4 ta kartochka
    _s(10, 15, "s10_t1", "sarlavha", 9, "10-slayd sarlavhasi 1-qator (1 so'z)"),
    _s(10, 23, "s10_t2", "sarlavha", 9, "10-slayd sarlavhasi 2-qator (1 so'z)"),
    _s(10, 24, "s10_k1", "sarlavha", 14, "1-kartochka nomi"),
    _s(10, 19, "s10_k1_m", "matn", 95, "1-kartochka izohi"),
    _s(10, 40, "s10_k2", "sarlavha", 14, "2-kartochka nomi"),
    _s(10, 31, "s10_k2_m", "matn", 95, "2-kartochka izohi"),
    _s(10, 32, "s10_k3", "sarlavha", 16, "3-kartochka nomi"),
    _s(10, 39, "s10_k3_m", "matn", 95, "3-kartochka izohi"),
    _s(10, 48, "s10_k4", "sarlavha", 16, "4-kartochka nomi"),
    _s(10, 47, "s10_k4_m", "matn", 95, "4-kartochka izohi"),
    _tozala(10, 20), _tozala(10, 21), _tozala(10, 22),

    # 11: 3 ta kartochka
    _s(11, 15, "s11_t1", "sarlavha", 9, "11-slayd sarlavhasi 1-qator (1 so'z)"),
    _s(11, 19, "s11_t2", "sarlavha", 9, "11-slayd sarlavhasi 2-qator (1 so'z)"),
    _s(11, 24, "s11_k1", "sarlavha", 18, "1-kartochka nomi"),
    _s(11, 23, "s11_k1_m", "matn", 95, "1-kartochka izohi"),
    _s(11, 33, "s11_k2", "sarlavha", 12, "2-kartochka nomi"),
    _s(11, 32, "s11_k2_m", "matn", 95, "2-kartochka izohi"),
    _s(11, 42, "s11_k3", "sarlavha", 16, "3-kartochka nomi"),
    _s(11, 41, "s11_k3_m", "matn", 95, "3-kartochka izohi"),
    _sobit(11, 50, "BATAFSIL"), _sobit(11, 54, "BATAFSIL"), _sobit(11, 58, "BATAFSIL"),
    _tozala(11, 16), _tozala(11, 17), _tozala(11, 18),

    # 12: ikki bo'lim
    _s(12, 15, "s12_t1", "sarlavha", 8, "12-slayd sarlavhasi 1-qator (1 so'z)"),
    _s(12, 21, "s12_t2", "sarlavha", 10, "12-slayd sarlavhasi 2-qator (1 so'z)"),
    _s(12, 16, "s12_sub", "sarlavha", 15, "Kichik sarlavha (2-3 so'z)"),
    _s(12, 17, "s12_a", "matn", 150, "1-bo'lim matni"),
    _s(12, 22, "s12_b", "matn", 150, "2-bo'lim matni"),
    _tozala(12, 18), _tozala(12, 19), _tozala(12, 20),

    # 13: ikki bo'lim
    _s(13, 18, "s13_t1", "sarlavha", 8, "13-slayd sarlavhasi 1-qator (1 so'z)"),
    _s(13, 23, "s13_t2", "sarlavha", 7, "13-slayd sarlavhasi 2-qator (1 so'z)"),
    _s(13, 19, "s13_a", "matn", 150, "1-bo'lim matni"),
    _s(13, 24, "s13_b", "matn", 150, "2-bo'lim matni"),
    _tozala(13, 20), _tozala(13, 21), _tozala(13, 22),

    # 14: ikki band
    _s(14, 15, "s14_t1", "sarlavha", 13, "14-slayd sarlavhasi 1-qator (masalan: kelajagi)"),
    _s(14, 19, "s14_t2", "sarlavha", 8, "14-slayd sarlavhasi 2-qator (1 so'z)"),
    _s(14, 27, "s14_k1", "sarlavha", 16, "1-band nomi"),
    _s(14, 28, "s14_k1_m", "matn", 150, "1-band izohi"),
    _s(14, 36, "s14_k2", "sarlavha", 12, "2-band nomi"),
    _s(14, 37, "s14_k2_m", "matn", 150, "2-band izohi"),
    _tozala(14, 16), _tozala(14, 17), _tozala(14, 18),

    # 15: yakun
    _sobit(15, 12, "E'TIBORINGIZ UCHUN"),
    _sobit(15, 13, "RAHMAT"),
    _s(15, 14, "yakun_izoh", "matn", 42, "Yakuniy slayddagi bir qatorli qisqa shior"),
    _sobit(15, 18, "KO'RISHGUNCHA"),
    _tozala(15, 15), _tozala(15, 16), _tozala(15, 17),
]

SHABLONLAR = {
    "moda": {
        "nomi": "Moda va dizayn",
        "fayl": os.path.join(PAPKA, "shablonlar", "moda.pptx"),
        "slotlar": MODA_SLOTLAR,
    },
}

# Shrift kengligini baholash (katta harflar uchun o'rtacha koeffitsiyent)
KENGLIK_KOEFF = 0.78


def _birinchi_run(shape):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    return p0.runs[0]


def _matnni_qoy(shape, matn):
    """Birinchi run'ning formatini saqlab, qolgan run va abzatslarni olib tashlaydi."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    run = _birinchi_run(shape)
    run.text = matn
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    return run


def _qirqish(matn, limit):
    """Matn limitdan uzun bo'lsa, so'z yoki gap chegarasida qirqadi."""
    matn = " ".join(str(matn).split())
    if limit <= 0 or len(matn) <= limit:
        return matn
    kesim = matn[:limit]
    nuqta = max(kesim.rfind(". "), kesim.rfind("! "), kesim.rfind("? "))
    if nuqta > limit * 0.6:
        return kesim[:nuqta + 1]
    bosh = kesim.rfind(" ")
    return (kesim[:bosh] if bosh > 0 else kesim).rstrip(",;:-") + "."


def _qirqish_sarlavha(matn, limit):
    """Sarlavhalar uchun: limitdan biroz uzun bo'lsa ham qoldiriladi (shrift
    kichrayadi), juda uzun bo'lsa so'z chegarasida qirqiladi, nuqta qo'yilmaydi."""
    matn = " ".join(str(matn).split()).rstrip(".")
    chegara = int(limit * 1.5)
    if len(matn) <= chegara:
        return matn
    kesim = matn[:chegara]
    bosh = kesim.rfind(" ")
    return (kesim[:bosh] if bosh > 0 else kesim).rstrip(",;:-")


def malumot_tekshir(kontent, slotlar):
    """AI javobidagi har bir kalitni tekshiradi: yo'qlarini bo'sh qoldiradi,
    uzunlarini qirqadi."""
    toza = {}
    for s in slotlar:
        k = s["kalit"]
        if not k:
            continue
        qiymat = kontent.get(k, "")
        if not isinstance(qiymat, str):
            qiymat = str(qiymat)
        if s["rejim"] == "sarlavha":
            toza[k] = _qirqish_sarlavha(qiymat, s["limit"])
        else:
            toza[k] = _qirqish(qiymat, s["limit"])
    return toza


def prompt_yaratish(shablon_kaliti, mavzu):
    slotlar = SHABLONLAR[shablon_kaliti]["slotlar"]
    qatorlar = []
    for s in slotlar:
        if s["kalit"]:
            qatorlar.append(
                f'- "{s["kalit"]}": {s["tavsif"]} (eng ko\'pi bilan {s["limit"]} belgi)'
            )
    return (
        f"Quyidagi mavzu bo'yicha taqdimot matnini o'zbek tilida (lotin yozuvi) yozing. "
        f"Mavzu: {mavzu}\n\n"
        "Faqat JSON obyekt qaytaring, boshqa hech narsa yozmang. Quyidagi kalitlarning "
        "har biri uchun matn bering. HAR BIR MATN BELGILANGAN BELGI LIMITIDAN OSHMASIN, "
        "chunki matn tayyor dizayndagi kichik joyga sig'ishi kerak. Sarlavhalar juda "
        "qisqa, 1-2 so'zdan iborat bo'lsin. Apostrof uchun oddiy ' belgisidan foydalaning "
        "(masalan: o'zbek, ma'no). Matnlar bir-biriga mantiqan bog'langan va faktik "
        "jihatdan ehtiyotkor bo'lsin, aniq bilmagan sana yoki raqamlarni to'qimang.\n\n"
        "Kalitlar:\n" + "\n".join(qatorlar)
    )


def _slaydda_shape(prs, slayd_no, shape_id):
    for sh in prs.slides[slayd_no - 1].shapes:
        if sh.shape_id == shape_id:
            return sh
    return None


def shablondan_yaratish(shablon_kaliti, kontent):
    """Shablonni ochib, kontent bilan to'ldiradi va BytesIO qaytaradi."""
    sozlama = SHABLONLAR[shablon_kaliti]
    slotlar = sozlama["slotlar"]
    kontent = malumot_tekshir(kontent, slotlar)
    prs = Presentation(sozlama["fayl"])

    for s in slotlar:
        sh = _slaydda_shape(prs, s["slayd"], s["id"])
        if sh is None or not sh.has_text_frame:
            continue
        rejim = s["rejim"]

        if rejim == "tozala":
            _matnni_qoy(sh, "")
            continue
        if rejim == "sobit":
            matn = s["matn"]
            asl_razmer = _birinchi_run(sh).font.size
            eni = sh.width
            run = _matnni_qoy(sh, matn)
            _sig_dirish(run, matn, asl_razmer, eni)
            continue

        matn = kontent.get(s["kalit"], "")
        if rejim == "sarlavha":
            asl_razmer = _birinchi_run(sh).font.size
            eni = sh.width
            matn = matn.upper()
            run = _matnni_qoy(sh, matn)
            _sig_dirish(run, matn, asl_razmer, eni)
        elif rejim == "katta":
            _matnni_qoy(sh, matn.upper())
        else:
            _matnni_qoy(sh, matn)

    fayl = io.BytesIO()
    prs.save(fayl)
    fayl.seek(0)
    return fayl


def _sig_dirish(run, matn, asl_razmer, eni_emu):
    """Bir qatorli sarlavha shablondagi eni ichiga sig'ishi uchun shriftni
    kerak bo'lsa kichraytiradi."""
    if not asl_razmer or not matn:
        return
    pt = asl_razmer.pt
    eni_pt = eni_emu / 12700
    kerakli = len(matn) * KENGLIK_KOEFF * pt
    if kerakli > eni_pt:
        pt = max(10, pt * eni_pt / kerakli)
    run.font.size = Pt(round(pt, 1))
